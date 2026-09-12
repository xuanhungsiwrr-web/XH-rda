import importlib.util
import json
import os
from pathlib import Path
import sys
import tempfile
import unittest
from unittest.mock import patch
import urllib.error
import zipfile

sys.path.insert(0,str(Path(__file__).resolve().parents[1]/'scripts'))
from xh_core import Project, relative, library
from xh_artifacts import edit_guard, retrieve, attachment, render, export_project
from xh_web import clean_html

class WorkspaceTests(unittest.TestCase):
    def setUp(self):
        self.tmp=tempfile.TemporaryDirectory()
        self.root=Path(self.tmp.name)/'project'
        Project.init(self.root, {k:'demo' for k in ['project_name','consultant','investor',
                     'report_type','design_stage','location']})
        self.p=Project(self.root)
        self.lib=Path(self.tmp.name)/'library'; self.lib.mkdir()
        (self.lib/'unknown-report.md').write_text('# Required\n\nMinimum A\n\nMinimum B',encoding='utf-8')
        self.p.select_requirements(self.lib,['unknown-report.md'])
        ids=[r['id'] for r in json.loads(self.p.read('requirements'))]
        self.p.define_sections([{'id':'A','title':'First','requirements':ids,'depends_on':[]},
            {'id':'B','title':'Dependent','requirements':[],'depends_on':['A']},
            {'id':'C','title':'Independent','requirements':[],'depends_on':[]}])
        self.p.approve('outline',self.p.head('outline')['id'],'reviewer')

    def tearDown(self):
        self.p.close(); self.tmp.cleanup()

    def write_all(self):
        for sid in ['A','B','C']: self.p.section(sid,'# '+sid+'\n\nContent '+sid)

    def test_no_fixed_report_type_and_coverage(self):
        self.assertEqual(len(library(self.lib)),1)
        with self.assertRaises(ValueError):
            self.p.define_sections([{'id':'New','title':'New','requirements':[]}])

    def test_incremental_invalidation_survives_resume(self):
        self.write_all(); old_c=self.p.head('section:C')['id']
        self.p.section('A','# A\n\nChanged',origin='human')
        self.p.close(); self.p=Project(self.root)
        self.assertTrue(self.p.stale('section:B'))
        self.assertFalse(self.p.stale('section:C'))
        self.assertEqual(self.p.head('section:C')['id'],old_c)
        self.assertEqual(self.p.plan()['batches'],[['B']])

    def test_parallel_respects_graph(self):
        self.assertEqual(self.p.plan('all',parallel=2)['batches'],[['A','C'],['B']])

    def test_human_edits_not_overwritten(self):
        self.p.section('A','# A\nOriginal')
        (self.root/'drafts/A.md').write_text('User changed',encoding='utf-8')
        with self.assertRaises(ValueError): self.p.section('A','AI overwrite')
        self.assertTrue(self.p.stale('section:A'))

    def test_revision_compare_and_swap(self):
        self.p.section('A','First'); old=self.p.head('section:A')['id']
        self.p.section('A','Human',origin='human')
        with self.assertRaises(ValueError): self.p.section('A','Late worker',expected=old)

    def test_conflict_preserved_until_human_resolves(self):
        r={'value':2,'unit':'m','source':'source#p1','quote':'2m','data_type':'PROJECT'}
        self.p.fact('height',r); self.p.resolve_fact('height',0,'user')
        self.p.section('A','# A\n{{fact:height}}')
        self.p.fact('height',{**r,'value':3,'quote':'3m'})
        current=json.loads(self.p.read('fact:height'))
        self.assertEqual(len(current['candidates']),2)
        self.assertTrue(self.p.stale('section:A'))
        self.assertFalse(self.p.approved('fact:height'))

    def test_approval_does_not_follow_edit(self):
        self.p.section('A','First'); h=self.p.head('section:A')
        self.p.approve('section:A',h['id'],'user')
        self.p.section('A','Second',origin='human')
        self.assertFalse(self.p.approved('section:A'))
        with self.assertRaises(ValueError): self.p.approve('section:A',h['id'],'user')

    def test_final_requires_reviews_and_human(self):
        self.write_all()
        with self.assertRaises(ValueError): self.p.assemble(True)
        for sid in ['A','B','C']:
            self.p.review(sid,[],'independent-reviewer',True)
            self.p.approve('section:'+sid,self.p.head('section:'+sid)['id'],'user')
        self.p.assemble(True)
        self.assertIn('Content C',self.p.read('report').decode())

    def test_rollback_preserves_old_dependency(self):
        self.write_all(); old=self.p.head('section:B')['id']
        self.p.section('A','Change',origin='human')
        result=self.p.rollback(old,'user')
        self.assertTrue(result['stale'])

    def test_path_escape(self):
        for name in ['../escape','D:\\escape','file.txt:stream']:
            with self.assertRaises(ValueError): relative(self.root,name)

    def test_cycle_rejected(self):
        self.p.put('x','sources/x.txt','X')
        self.p.put('y','sources/y.txt','Y',deps=['x'])
        with self.assertRaises(ValueError): self.p.put('x','sources/x.txt','Z',deps=['y'])

    def test_worker_cannot_claim_new_input_revision(self):
        self.p.put('source:x','sources/x.md','Old')
        old=self.p.head('source:x')['id']
        self.p.put('source:x','sources/x.md','New',origin='human')
        with self.assertRaises(ValueError):
            self.p.put('candidate:x','reviews/x.md','Output from old input',deps=['source:x'],
                       expected_deps={'source:x':old})

    def test_html_normalization_preserves_table(self):
        raw='<main><p>'+'Source evidence. '*20+'</p><table><tr><th>m</th></tr><tr><td>2</td></tr></table></main>'
        text,links=clean_html(raw,'https://example.com')
        self.assertIn('| 2 |',text)

    def test_humanizer_guard(self):
        self.assertFalse(edit_guard('Height {{fact:h}} 2m','Height 3m')['protected_content_preserved'])
        self.assertTrue(edit_guard('Height {{fact:h}}','Cao trình {{fact:h}}')['protected_content_preserved'])

    def test_attachment_and_archive(self):
        (self.root/'calculations/a.xlsx').write_bytes(b'fixture-binary')
        attachment(self.p,'calc','calculations/a.xlsx','Result from sheet1!B2')
        self.p.section('A','Result',deps=['attachment:calc'])
        out=Path(self.tmp.name)/'export.zip'; export_project(self.p,out)
        with zipfile.ZipFile(out) as z:
            self.assertIn('calculations/a.xlsx',z.namelist())
            self.assertIn('.xh/state.sqlite',z.namelist())

    def test_render_one_document_and_missing_template_style(self):
        import docx
        self.write_all()
        d=docx.Document(); d.save(self.root/'templates/basic.docx')
        styles={k:'Normal' for k in ['than_bai','bullet','so_thu_tu','o_tieu_de','o_noi_dung',
                'caption_bang','doan_chua_anh','caption_hinh']}
        styles.update(heading=['Heading '+str(i) for i in range(1,7)],bang='Table Grid')
        (self.root/'templates/contract.json').write_text(json.dumps({'styles':styles}),encoding='utf-8')
        result=render(self.p,'templates/basic.docx','templates/contract.json')
        rendered=docx.Document(self.root/result['path'])
        text='\n'.join(p.text for p in rendered.paragraphs)
        self.assertTrue(all('Content '+s in text for s in ['A','B','C']))
        styles['than_bai']='MissingStyle'
        (self.root/'templates/contract.json').write_text(json.dumps({'styles':styles}),encoding='utf-8')
        with self.assertRaises(ValueError): render(self.p,'templates/basic.docx','templates/contract.json')

if __name__=='__main__': unittest.main()
