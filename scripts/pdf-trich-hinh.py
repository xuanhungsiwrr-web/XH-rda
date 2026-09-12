#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
pdf-trich-hinh.py — trích hình và sơ đồ từ PDF kỹ thuật bằng docling.

Nguồn: skill `pdf-to-word-engineering` của plugin xh-tuvan-v2, sửa lại 25/08/2026.

HAI CHẾ ĐỘ:

  --che-do hinh   (MẶC ĐỊNH, đúng kiến trúc ba lớp)
      Trích ảnh ra thư mục + sinh một file .md danh mục hình.
      Bản thảo Markdown chèn hình bằng cú pháp ![](duong-dan), caption do
      render_report.py sinh bằng trường STYLEREF + SEQ. KHÔNG dựng .docx ở đây.

  --che-do word   (chỉ dùng để ĐỌC NHANH hồ sơ scan, không phải để giao nộp)
      Dựng một .docx thô gồm text + ảnh, như bản gốc của v2.
      File này là bản đọc tham khảo. Tuyệt đối không dùng làm bản thảo báo cáo —
      nó không đi qua khuôn KhuonMau_NamQuoc.dotx, không có mục lục, không có
      caption đánh số bằng field.

CÀI ĐẶT (docling nặng, kéo theo mô hình nhận diện — chỉ cài khi cần):
      pip install docling python-docx --break-system-packages

DÙNG:
      python3 pdf-trich-hinh.py "BanVe-PDF/CAU SONG DOC.pdf" --ra 00_input/Hinh/
      python3 pdf-trich-hinh.py hoso.pdf --che-do word --ra doc-nhanh.docx
"""

import argparse
import sys
from pathlib import Path


def _nap_docling():
    try:
        from docling.document_converter import DocumentConverter, PdfFormatOption
        from docling.datamodel.pipeline_options import PdfPipelineOptions
        from docling_core.types.doc import PictureItem, TextItem
    except ImportError:
        print("[LỖI] Chưa cài docling.\n"
              "      pip install docling --break-system-packages\n"
              "      (thư viện nặng, kéo theo mô hình nhận diện — cài khi thật cần)",
              file=sys.stderr)
        sys.exit(3)
    return DocumentConverter, PdfFormatOption, PdfPipelineOptions, PictureItem, TextItem


def doc_pdf(pdf_path: str):
    DocumentConverter, PdfFormatOption, PdfPipelineOptions, PictureItem, TextItem = _nap_docling()
    opts = PdfPipelineOptions()
    opts.generate_picture_images = True
    conv = DocumentConverter(format_options={"pdf": PdfFormatOption(pipeline_options=opts)})
    return conv.convert(str(pdf_path)).document, PictureItem, TextItem


def che_do_hinh(pdf_path: str, thu_muc_ra: str, tien_to: str) -> str:
    """Trích ảnh ra thư mục + sinh danh mục .md. Không dựng Word."""
    doc_data, PictureItem, TextItem = doc_pdf(pdf_path)
    ra = Path(thu_muc_ra)
    ra.mkdir(parents=True, exist_ok=True)

    dong = ["# Danh mục hình trích từ `%s`" % Path(pdf_path).name, "",
            "> Sinh bằng `pdf-trich-hinh.py` ngày %s. Caption dưới đây là **gợi ý**, "
            "phải sửa lại theo nội dung thật trước khi đưa vào bản thảo." % __import__("datetime").date.today(),
            "", "| File | Caption gợi ý | Chèn vào bản thảo |", "|---|---|---|"]

    n = 0
    for item, _lv in doc_data.iterate_items():
        if isinstance(item, PictureItem):
            n += 1
            ten = f"{tien_to}{n:02d}.png"
            with open(ra / ten, "wb") as fp:
                item.get_image(doc_data).save(fp, "PNG")
            cap = getattr(item, "caption", None) or ""
            cap = str(cap).strip() or "(chưa có caption — tự đặt)"
            dong.append(f"| `{ten}` | {cap} | `![{cap}]({ra.name}/{ten})` |")

    dong += ["", f"**Tổng: {n} hình.**", "",
             "Cách dùng: chép dòng cột cuối vào bản thảo `.md` tại đúng vị trí. "
             "`render_report.py` sẽ tự sinh số hiệu `Hình <chương>-<số>` bằng trường "
             "STYLEREF + SEQ — **không gõ số hiệu bằng tay**."]

    dm = ra / "_danh-muc-hinh.md"
    dm.write_text("\n".join(dong) + "\n", encoding="utf-8")
    return f"Đã trích {n} hình vào {ra}/ và ghi danh mục {dm}"


def che_do_word(pdf_path: str, ra_docx: str) -> str:
    """Dựng .docx thô để đọc nhanh. KHÔNG phải bản thảo báo cáo."""
    try:
        from docx import Document
        from docx.shared import Inches, Pt
        from docx.enum.text import WD_ALIGN_PARAGRAPH
    except ImportError:
        print("[LỖI] Chưa cài python-docx:  pip install python-docx --break-system-packages",
              file=sys.stderr)
        sys.exit(3)

    doc_data, PictureItem, TextItem = doc_pdf(pdf_path)
    out = Path(ra_docx)
    thu_muc_anh = out.parent / "extracted_images"
    thu_muc_anh.mkdir(parents=True, exist_ok=True)

    doc = Document()
    doc.add_heading("BẢN ĐỌC NHANH — KHÔNG PHẢI BẢN THẢO BÁO CÁO", level=0)
    doc.add_paragraph(
        "File này do pdf-trich-hinh.py dựng để đọc nhanh nội dung PDF. "
        "Nó không đi qua khuôn KhuonMau_NamQuoc.dotx, không có mục lục và không có "
        "caption đánh số bằng field. Không dùng để giao nộp."
    )

    dem = 1
    for item, _lv in doc_data.iterate_items():
        if isinstance(item, TextItem):
            t = (item.text or "").strip()
            if t:
                p = doc.add_paragraph(t)
                p.paragraph_format.space_after = Pt(6)
        elif isinstance(item, PictureItem):
            img = thu_muc_anh / f"figure_{dem}.png"
            with open(img, "wb") as fp:
                item.get_image(doc_data).save(fp, "PNG")
            p_img = doc.add_paragraph()
            p_img.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p_img.add_run().add_picture(str(img), width=Inches(5.5))
            cap = getattr(item, "caption", None) or f"Hình minh hoạ {dem}"
            p_cap = doc.add_paragraph(f"Hình {dem}: {cap}")
            p_cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p_cap.runs[0].font.italic = True
            p_cap.runs[0].font.size = Pt(10)
            dem += 1

    doc.save(str(out))
    return f"Đã dựng bản đọc nhanh: {out} ({dem - 1} hình)"


def main():
    ap = argparse.ArgumentParser(description="Trích hình từ PDF kỹ thuật bằng docling")
    ap.add_argument("pdf", help="File PDF đầu vào")
    ap.add_argument("--che-do", choices=["hinh", "word"], default="hinh",
                    help="hinh = trích ảnh + danh mục .md (mặc định); word = bản đọc nhanh .docx")
    ap.add_argument("--ra", required=True, help="Thư mục ra (chế độ hinh) hoặc file .docx (chế độ word)")
    ap.add_argument("--tien-to", default="H", help="Tiền tố tên file ảnh, mặc định 'H'")
    a = ap.parse_args()

    if not Path(a.pdf).exists():
        print(f"[LỖI] Không thấy file: {a.pdf}", file=sys.stderr)
        sys.exit(1)

    if a.che_do == "hinh":
        print(che_do_hinh(a.pdf, a.ra, a.tien_to))
    else:
        print(che_do_word(a.pdf, a.ra))


if __name__ == "__main__":
    main()
