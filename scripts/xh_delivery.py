"""Immutable report delivery pairs and revision-aware feedback import."""
from __future__ import annotations

import difflib
import json
import os
from pathlib import Path
import re

from xh_core import digest, encoded, relative, stamp
from xh_layout import PLUGIN_VERSION


def _safe_release(value):
    if not isinstance(value, str) or not re.fullmatch(r"[A-Za-z0-9][A-Za-z0-9._-]{0,119}", value):
        raise ValueError("release_id must use letters, numbers, dot, underscore or hyphen")
    return value


def _exclusive_seed(path, data):
    path = Path(path)
    path.parent.mkdir(parents=True, exist_ok=True)
    try:
        with open(path, "xb") as stream:
            stream.write(data); stream.flush(); os.fsync(stream.fileno())
        return True
    except FileExistsError:
        if digest(path.read_bytes()) != digest(data):
            raise ValueError("Delivery path exists with different content: " + str(path))
        return False


def _row(project, pair_id=None, release_id=None):
    if pair_id:
        row = project.db.execute("SELECT * FROM delivery_pairs WHERE pair_id=?", (pair_id,)).fetchone()
    else:
        row = project.db.execute("SELECT * FROM delivery_pairs WHERE release_id=?", (release_id,)).fetchone()
    return dict(row) if row else None


def _verify_baseline(project, row):
    output = relative(project.root, row["output_path"])
    if not output.is_file():
        raise ValueError("Delivery baseline is missing: " + row["output_path"])
    current = digest(output.read_bytes())
    if current != row["baseline_hash"]:
        raise ValueError("Delivery baseline changed outside the revision store; feedback import refused")
    head = project.head(row["output_artifact"])
    if not head or head["hash"] != row["baseline_hash"] or project.stale(row["output_artifact"]):
        raise ValueError("Delivery baseline revision is missing or stale")
    return output


def deliver_bytes(project, release_id, data, output_name=None, simulate_failure_after_output=False):
    """Create a byte-identical Outputs/Feedback pair without overwriting either path."""
    release_id = _safe_release(release_id)
    if not isinstance(data, bytes):
        raise ValueError("Delivery content must be bytes")
    output_name = output_name or release_id + ".docx"
    if Path(output_name).name != output_name or not output_name.lower().endswith(".docx"):
        raise ValueError("output_name must be a DOCX filename, not a path")
    output_path = project.path("outputs", output_name)
    feedback_name = Path(output_name).stem + "_XHedited.docx"
    feedback_path = project.path("feedback", feedback_name)
    baseline_hash = digest(data)
    pair_id = "PAIR-" + digest(encoded([release_id, output_name, baseline_hash]))[:24]
    output_artifact = "delivery-output:" + release_id
    feedback_artifact = "delivery-feedback:" + release_id
    row = _row(project, release_id=release_id)
    if row and (row["pair_id"] != pair_id or row["baseline_hash"] != baseline_hash):
        raise ValueError("release_id already belongs to a different immutable baseline")
    now = stamp()
    if not row:
        with project.db:
            project.db.execute("INSERT INTO delivery_pairs VALUES(?,?,?,?,?,?,?,?,?,?,?)", (
                pair_id, release_id, output_artifact, feedback_artifact, output_path, feedback_path,
                baseline_hash, "preparing", PLUGIN_VERSION, now, now))
        row = _row(project, pair_id=pair_id)

    output = relative(project.root, output_path)
    if row["status"] != "complete":
        _exclusive_seed(output, data)
        result = project.put(output_artifact, output_path, data, "system", ["report"],
                             {"pair_id": pair_id, "release_id": release_id, "immutable_baseline": True})
        with project.db:
            project.db.execute("UPDATE delivery_pairs SET status='output_created',updated=? WHERE pair_id=?",
                               (stamp(), pair_id))
        if simulate_failure_after_output:
            raise RuntimeError("Simulated failure after output creation")
    else:
        result = {"artifact": output_artifact, "revision": project.head(output_artifact)["id"], "unchanged": True}
    _verify_baseline(project, _row(project, pair_id=pair_id))

    feedback = relative(project.root, feedback_path)
    feedback_preexisting = feedback.exists()
    if not project.head(feedback_artifact):
        _exclusive_seed(feedback, data)
        project.put(feedback_artifact, feedback_path, data, "system", [output_artifact],
                    {"pair_id": pair_id, "baseline_hash": baseline_hash, "editable_copy": True})
    elif not feedback.is_file():
        # Resume only from the immutable baseline; never from an absent/unknown feedback copy.
        _exclusive_seed(feedback, data)
        project.put(feedback_artifact, feedback_path, data, "system", [output_artifact],
                    {"pair_id": pair_id, "baseline_hash": baseline_hash, "resumed": True})
    with project.db:
        project.db.execute("UPDATE delivery_pairs SET status='complete',updated=? WHERE pair_id=?", (stamp(), pair_id))
    feedback_hash = digest(feedback.read_bytes())
    return {**result, "pair_id": pair_id, "release_id": release_id,
            "output_path": output_path, "feedback_path": feedback_path,
            "sha256": baseline_hash, "byte_identical": feedback_hash == baseline_hash,
            "feedback_modified": feedback_hash != baseline_hash,
            "resumed": row["status"] != "preparing" or feedback_preexisting}


def delivery_status(project, pair_id=None, release_id=None):
    if not pair_id and not release_id:
        return [dict(r) for r in project.db.execute("SELECT * FROM delivery_pairs ORDER BY created")]
    row = _row(project, pair_id=pair_id, release_id=release_id)
    if not row:
        raise ValueError("Unknown delivery pair")
    output = _verify_baseline(project, row)
    feedback = relative(project.root, row["feedback_path"])
    return {**row, "output_sha256": digest(output.read_bytes()),
            "feedback_sha256": digest(feedback.read_bytes()) if feedback.is_file() else None,
            "feedback_exists": feedback.is_file()}


def _docx_blocks(path):
    try:
        import docx
        document = docx.Document(path)
    except Exception as exc:
        raise ValueError("Feedback DOCX cannot be parsed; preserve it and inspect manually") from exc
    blocks = [p.text.strip() for p in document.paragraphs if p.text.strip()]
    for table in document.tables:
        for row in table.rows:
            cells = [cell.text.strip() for cell in row.cells]
            if any(cells):
                blocks.append(" | ".join(cells))
    return blocks


def import_feedback(project, pair_id=None, release_id=None):
    """Import exactly one edited copy after proving its immutable baseline pair."""
    row = _row(project, pair_id=pair_id, release_id=release_id)
    if not row:
        raise ValueError("Unknown delivery pair")
    baseline = _verify_baseline(project, row)
    feedback = relative(project.root, row["feedback_path"])
    if not feedback.is_file():
        raise ValueError("Feedback copy is missing")
    raw = feedback.read_bytes(); feedback_hash = digest(raw)
    if feedback_hash == row["baseline_hash"]:
        return {"pair_id": row["pair_id"], "changed": False, "imported": False,
                "reason": "feedback remains byte-identical to baseline"}
    old = project.db.execute("SELECT * FROM feedback_imports WHERE pair_id=? AND feedback_hash=?",
                             (row["pair_id"], feedback_hash)).fetchone()
    if old:
        return {"pair_id": row["pair_id"], "changed": True, "imported": False,
                "duplicate": True, "feedback_revision": old["feedback_revision"],
                "analysis_artifact": old["analysis_artifact"]}

    expected = project.head(row["feedback_artifact"])["id"]
    imported = project.put(row["feedback_artifact"], row["feedback_path"], raw, "human",
                           [row["output_artifact"]], {"pair_id": row["pair_id"],
                           "baseline_hash": row["baseline_hash"], "feedback_hash": feedback_hash},
                           actor="user-feedback", expected=expected)
    before_blocks = _docx_blocks(baseline)
    after_blocks = _docx_blocks(feedback)
    diff = list(difflib.unified_diff(before_blocks, after_blocks, fromfile="baseline", tofile="feedback", lineterm=""))
    classification = "formatting-only" if before_blocks == after_blocks else "content-and-or-structure"
    analysis = {"pair_id": row["pair_id"], "release_id": row["release_id"],
                "baseline_hash": row["baseline_hash"], "feedback_hash": feedback_hash,
                "feedback_revision": imported["revision"], "classification": classification,
                "diff": diff, "candidate_status": "not-created",
                "note": "A raw DOCX/XML difference is not treated as a reusable lesson."}
    analysis_id = "feedback-analysis:" + row["pair_id"] + ":" + feedback_hash[:12]
    analysis_path = project.path("edit_analysis", row["release_id"] + "-" + feedback_hash[:12] + ".json")
    saved = project.put(analysis_id, analysis_path, encoded(analysis), "system",
                        [row["output_artifact"], row["feedback_artifact"]])
    with project.db:
        project.db.execute("INSERT INTO feedback_imports VALUES(?,?,?,?,?)", (
            row["pair_id"], feedback_hash, imported["revision"], analysis_id, stamp()))
    return {"pair_id": row["pair_id"], "changed": True, "imported": True,
            "feedback_revision": imported["revision"], "analysis_artifact": analysis_id,
            "analysis_path": saved["path"], "classification": classification,
            "learning_candidate_created": False}
